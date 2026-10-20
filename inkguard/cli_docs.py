"""The three documents a run writes beside the redacted files, in one house
style: the same header, the same numbered blocks, the
same footer, kept compact.

    4_Audit_trail.pdf        the record of the run, with every file by id
    5_Citations.txt          references to copy, in plain text so nothing
                             unwanted comes along with the paste
    6_PLEASE-RETURN-...docx  the feedback report: the user's own box at the
                             top, the software's record below, to be emailed

Each is built from the run's report (a dict, see cli_report.build_report),
so the three always agree with each other and with the machine copy in the
hidden folder. Nothing from the contents of the user's files goes into any
of them; the audit trail does show the input folder's name and the
computer's name.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

LOGO = Path(__file__).with_name("logo.png")

# The house colours, the same as the window's.
INK = "#1f2328"
MUTED = "#59636e"
RULE = "#d0d7de"
SUBTLE = "#f6f8fa"
ACCENT = "#0969da"
RED = "#cf222e"
AMBER = "#9a6700"


def _when(utc: str | None) -> str:
    return (utc or "").replace("T", " ").replace("Z", " UTC")


# =====================================================================
# 4_Audit_trail.pdf
# =====================================================================

#: The folder each outcome word (cli_report._OUTCOME_WORDS) lands in.
_FOLDER_OF = {"redacted, no flags": "1_REDACTED-NO-FLAGS", "redacted, flagged": "2_REDACTED-FLAGGED",
              "not processed": "3_NOT-PROCESSED-MAY-CONTAIN-PHI",
              "not processed (system file)": "3_NOT-PROCESSED-MAY-CONTAIN-PHI (system file)"}


def audit_pdf(report: dict, path: Path) -> None:
    """Write the audit trail. Fixed blocks, in this order: the run, the
    results, the checks; then, as long as the run needs, files by folder and
    what was removed and flagged; then an appendix of every file by id. A
    footer on every page."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    from inkguard import cli_report as R

    c = colors.HexColor
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=7.6, leading=9.6, textColor=c(INK))
    small = ParagraphStyle("small", parent=base, fontSize=6.8, leading=8.4, textColor=c(MUTED))
    title = ParagraphStyle("title", parent=base, fontName="Helvetica-Bold", fontSize=14, leading=17)
    meta = ParagraphStyle("meta", parent=small, alignment=TA_RIGHT)
    head = ParagraphStyle("head", parent=base, fontName="Helvetica-Bold", fontSize=7.4, leading=9,
                          textColor=c(ACCENT), spaceBefore=7, spaceAfter=2)
    mono = ParagraphStyle("mono", parent=base, fontName="Courier", fontSize=6.8, leading=8.4)

    def P(text, style=base):
        return Paragraph(text, style)

    def esc(t) -> str:
        return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    def grid(rows, widths, header=False, zebra=True):
        t = Table(rows, colWidths=widths, hAlign="LEFT")
        st = [("FONT", (0, 0), (-1, -1), "Helvetica", 7.4), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6),
              ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
              ("LINEBELOW", (0, 0), (-1, -1), 0.25, c(RULE)), ("BOX", (0, 0), (-1, -1), 0.5, c(RULE))]
        if header:
            st += [("BACKGROUND", (0, 0), (-1, 0), c(SUBTLE)), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.2)]
        if zebra:
            for i in range(1 if header else 0, len(rows)):
                if i % 2 == (0 if header else 1):
                    st.append(("BACKGROUND", (0, i), (-1, i), c("#fbfcfd")))
        t.setStyle(TableStyle(st))
        return t

    run = report.get("run_id") or "none"
    sw = report.get("software") or {}
    rr = report.get("run_record") or {}
    at, pc = rr.get("entered") or {}, rr.get("computer") or {}
    pl = report.get("piles") or {}
    st = report.get("status") or {}
    fs = report.get("files") or {}
    loc = report.get("locations") or {}
    W = LETTER[0] - 1.0 * inch
    given = (lambda k: esc(at.get(k)) if at.get(k) else '<font color="#8c959f">not given</font>')

    story = []
    head_cell = P(f"Inkguard · {R.AUDIT_NAME}", title)
    # The logo as vector shapes: sharp at any zoom.
    from reportlab.graphics.shapes import Circle, Drawing, Rect
    mark = Drawing(13, 13)
    mark.add(Circle(6.5, 6.5, 5.85, fillColor=c("#A91636"), strokeColor=None))
    mark.add(Rect(2.08, 4.55, 8.84, 3.9, rx=1.95, ry=1.95, fillColor=c("#000000"), strokeColor=None))
    head_cell = Table([[mark, head_cell]], colWidths=[18, W * 0.6 - 18],
                      style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)])
    story.append(Table([[head_cell,
                         P(f"<b>Run</b> {esc(run)} · <b>Version</b> {esc(sw.get('version', R._version()))} · "
                           f"<b>Written</b> {esc(_when(report.get('generated_utc')))}", meta)]],
                       colWidths=[W * 0.6, W * 0.4], style=[("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                                                            ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                                            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                                            ("LINEBELOW", (0, 0), (-1, 0), 1.2, c(INK))]))
    story.append(Spacer(1, 3))
    story.append(P(" ".join(R.AUDIT_STATEMENT) + " " + esc(R._verify()), small))

    # 1. Run
    story.append(P("1  RUN", head))
    kv = [
        ("Run by", given("operator") + (f", {esc(at['operator_role'])}" if at.get("operator_role") else ""),
         "Started", esc(_when(pc.get("started_utc")))),
        ("Principal investigator", given("pi"), "Finished", esc(_when(rr.get("finished_utc")))),
        ("Institution", given("institution"), "Computer", esc(pc.get("computer_name") or "")),
        ("IRB or ethics approval", given("protocol"), "System",
         esc(f"{pc.get('operating_system', '')} · Python {pc.get('python', '')}")),
        ("Profile", esc(R._profile_fact(report)).replace("**", "").replace("safe_harbor (45 CFR 164.514(b)(2))",
         "Safe Harbor identifier list (45 CFR 164.514(b)(2)(i)); not a de-identification determination")
         .replace("safe_harbor (", "Safe Harbor identifier list (")
         .replace("limited_data_set (", "Limited Data Set (").replace("safe_harbor with", "HIPAA Safe Harbor with"),
         "Time zone", esc(pc.get("time_zone") or "")),
        ("File and folder names", "randomised" if report.get("scramble") else "kept unless they held an identifier",
         "Names in notes", "replaced with [NAME]" if report.get("names", {}).get("redaction_mode") == "mask"
         else "replaced with a consistent code"),
    ]
    rows = [[P(f'<font color="{MUTED}">{a}</font>'), P(b), P(f'<font color="{MUTED}">{x}</font>'), P(y)]
            for a, b, x, y in kv]
    story.append(grid(rows, [W * 0.17, W * 0.35, W * 0.12, W * 0.36]))
    story.append(P("Details for the person and study are recorded as entered and are not verified.", small))

    # 2. Results
    story.append(P("2  RESULTS", head))
    res = [[P("<b>Folder</b>"), P("<b>Files</b>"), P("<b>Meaning</b>")],
           [P(f'<font color="{ACCENT}"><b>{R.PILE_DONE}</b></font>'), P(f"{pl.get('done', 0):,}"),
            P("Redacted. Pending human review.")],
           [P(f'<font color="{AMBER}"><b>{R.PILE_FLAGGED}</b></font>'), P(f"{pl.get('flagged', 0):,}"),
            P(f"Redacted. Words flagged ({R.WORDS_FILE} in each patient folder). Pending human review.")],
           [P(f'<font color="{RED}"><b>{R.PILE_NOT}</b></font>'), P(f"{pl.get('not_processed', 0):,}"),
            P("Not processed: exact copies of the originals. Not redacted.")],
           [P("<b>Received</b>"), P(f"<b>{pl.get('received', 0):,}</b>"), P("")]]
    story.append(grid(res, [W * 0.33, W * 0.08, W * 0.59], header=True, zebra=False))
    story.append(Spacer(1, 2))
    story.append(grid([[P(f'<font color="{MUTED}">Input folder</font>'), P(esc(loc.get("input", "")), mono)],
                       [P(f'<font color="{MUTED}">Results folder</font>'), P(esc(loc.get("results", "")), mono)]],
                      [W * 0.17, W * 0.83], zebra=False))

    # 3. Checks
    ok = pl.get("every_file_in_exactly_one")

    def mark(v):
        return ('<font color="#1a7f37"><b>Yes</b></font>' if v else
                f'<font color="{RED}"><b>No</b></font>' if v is False else "Not checked")
    story.append(P("3  CHECKS", head))
    checks = [[P("Every file received is in exactly one folder"), P(mark(ok)),
               P("Code fingerprint (SHA-256)"), P(esc(sw.get("code_fingerprint_sha256", "")), mono)],
              [P("Originals unchanged after the run (SHA-256)"),
               P(mark(st.get("originals_unchanged_checked_by_sha256"))),
               P("Input set fingerprint (SHA-256)"), P(esc(report.get("input_set_fingerprint") or ""), mono)],
              [P("Stopped as a precaution / refused at intake"),
               P(f"{st.get('of_which_stopped_as_a_precaution', 0):,} / {fs.get('intake_guard_refused', 0):,}"),
               P("Fingerprints of every file written"), P(f"{R.TECH_DIR}/{R.MANIFEST_FILE}", mono)]]
    story.append(grid(checks, [W * 0.27, W * 0.07, W * 0.2, W * 0.46], zebra=False))
    for prob in (fs.get("problems") or [])[:10]:
        story.append(P(f'<font color="{RED}">Problem: {esc(prob)}</font>', small))

    # 4. Files by folder
    folders = report.get("by_folder") or []
    if folders:
        story.append(P("4  FILES BY FOLDER", head))
        cap = 60
        rows = [[P("<b>Folder</b>"), P("<b>Files</b>"), P("<b>No flags</b>"), P("<b>Flagged</b>"),
                 P("<b>Not processed</b>")]]
        for f in folders[:cap]:
            rows.append([P(esc(f["folder"]), mono), P(f"{f['files']:,}"), P(f"{f['done']:,}"),
                         P(f"{f['flagged']:,}"), P(f"{f['not_processed']:,}")])
        story.append(grid(rows, [W * 0.44, W * 0.12, W * 0.14, W * 0.14, W * 0.16], header=True))
        if len(folders) > cap:
            story.append(P(f"{len(folders) - cap:,} more folders. Every file, by id, is listed in the "
                           f"appendix.", small))
        else:
            story.append(P(f"Every file, by id, is listed in the appendix. Original names: "
                           f"{R.KEY_FILE} in {R.PILE_NOT}.", small))

    # 5. Removed, flagged, other events
    story.append(P("5  REMOVED, FLAGGED, OTHER EVENTS", head))
    ids = report.get("identifiers_removed") or {}
    reasons = (report.get("review_check") or {}).get("reasons") or {}
    dc = report.get("documents") or {}
    lines = [("Removed", " · ".join(f"{esc(k)} {v:,}" for k, v in ids.items()) or "none"),
             ("Flagged words", " · ".join(f"{esc(k.replace('_', ' '))} {v:,}" for k, v in reasons.items())
              or "none")]
    if dc:
        lines.append(("Hidden content", " · ".join(f"{esc(k)} {v}" for k, v in
                                                   dc.get("hidden_content_found", {}).items()) or "none"))
    if fs.get("intake_guard_notes"):
        lines.append(("Intake check", " · ".join(f"{esc(k)} {v}" for k, v in fs["intake_guard_notes"].items())))
    d = report.get("deidentification") or {}
    if d.get("output_is_still_phi"):
        lines.append(("Limited Data Set", f"output remains PHI; towns removed to be safe "
                                          f"{d.get('towns_removed_to_be_safe', 0):,}; names or IDs inside "
                                          f"kept dates or towns {d.get('names_or_ids_removed_inside_kept_dates_or_towns', 0):,}"))
    codes = report.get("support_codes") or {}
    lines.append(("Support codes", " · ".join(f"{esc(k)} ×{v}" for k, v in codes.items()) or "none"))
    story.append(grid([[P(f'<font color="{MUTED}">{a}</font>'), P(b)] for a, b in lines],
                      [W * 0.17, W * 0.83], zebra=False))
    # Appendix: every file by id. Ids, folders and fingerprints only; never a
    # file or folder name, which could carry an identifier the tool missed.
    log = report.get("file_log") or []
    if log:
        story.append(PageBreak())
        story.append(P("APPENDIX  FILES BY ID", head))
        story.append(P("Each file received, by id: the folder it went to, and the first 16 "
                       "characters of the SHA-256 fingerprint of the copy written. Files in "
                       f"{R.PILE_NOT} are exact copies and have no fingerprint here. Full "
                       f"fingerprints: {R.TECH_DIR}/{R.MANIFEST_FILE}.", small))
        rows = [[P("<b>File id</b>"), P("<b>Folder</b>"), P("<b>Fingerprint</b>")]]
        for f in log:
            rows.append([P(esc(f.get("id", "")), mono), P(esc(_FOLDER_OF.get(f.get("outcome", ""),
                                                                          f.get("outcome", "")))),
                         P(esc((f.get("sha256") or "—")[:16]), mono)])
        story.append(grid(rows, [W * 0.22, W * 0.5, W * 0.28], header=True))
    story.append(Spacer(1, 6))
    story.append(P(f"<i>End of {R.AUDIT_NAME.lower()}.</i>", small))

    about = (f"About this document. Written by Inkguard {R._version()}, code fingerprint "
             f"{R._fingerprint()[:16]}, for run {run}.")
    licence = f"Open-source software · Apache License 2.0 · {R.SOURCE.replace('https://', '')}"

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(c(RULE))
        canvas.setLineWidth(0.5)
        canvas.line(0.5 * inch, 0.5 * inch, LETTER[0] - 0.5 * inch, 0.5 * inch)
        canvas.setFont("Helvetica", 6.3)
        canvas.setFillColor(c(MUTED))
        canvas.drawString(0.5 * inch, 0.37 * inch, about)
        canvas.drawString(0.5 * inch, 0.26 * inch, licence)
        canvas.drawRightString(LETTER[0] - 0.5 * inch, 0.26 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                            topMargin=0.45 * inch, bottomMargin=0.65 * inch,
                            title=f"Inkguard audit trail, run {run}", author="Inkguard", subject=R.AUDIT_NAME,
                            creator=f"Inkguard {R._version()}", invariant=1)
    doc.build([KeepTogether(story[:3])] + story[3:], onFirstPage=footer, onLaterPages=footer)


# =====================================================================
# 5_Citations.txt
# =====================================================================

def cite_txt(report: dict) -> str:
    """Plain text so a reference pastes cleanly. Version and date on top;
    five separated blocks; the same footer as the other documents."""
    from inkguard import cli_report as R
    sw = report.get("software") or {"version": R._version(), "code_fingerprint_sha256": R._fingerprint()}
    version, fp = sw["version"], sw["code_fingerprint_sha256"]
    date = (report.get("generated_utc") or "")[:10]
    year = date[:4]
    from inkguard.fluorite import DOI
    bar, rule = "=" * 72, "-" * 72

    def block(n, name, body):
        return ["", f"[{n}] {name}", rule, *body]
    doi_url = f"https://doi.org/{DOI}" if DOI else R.SOURCE
    pub = "Zenodo" if DOI else "GitHub"
    return "\n".join([
        f"INKGUARD · {R.CITE_NAME.upper()}",
        f"Run {report.get('run_id', '')} · Version {version} · Written {date} UTC"
        + (f" · DOI {DOI}" if DOI else ""),
        bar,
        *block(1, "APA", [f"Wajahat, A. ({year}). Inkguard (Version {version}) [Computer software]. {pub}. {doi_url}"]),
        *block(2, "AMA", [f"Wajahat A. Inkguard. Version {version}. {pub}; {year}. "
                          + (f"doi:{DOI}" if DOI else f"Accessed [date]. {R.SOURCE}")]),
        *block(3, "Vancouver", [f"Wajahat A. Inkguard [computer program]. Version {version}. "
                                f"{'Geneva: Zenodo' if DOI else 'GitHub'}; {year} [cited YYYY Mon DD]. "
                                f"Available from: {doi_url}" + (f" doi: {DOI}" if DOI else "")]),
        *block(4, "BibTeX", [f"@software{{wajahat_inkguard_{year},", "  author    = {Wajahat, Alize},",
                             "  title     = {Inkguard},", f"  version   = {{{version}}},",
                             f"  year      = {{{year}}},", f"  publisher = {{{pub}}},",
                             *([f"  doi       = {{{DOI}}},"] if DOI else []), f"  url       = {{{doi_url}}}", "}"]),
        *block(5, "Methods paragraph (template; edit the parts in [brackets])",
               [report.get("methods_paragraph_draft") or ""]),
        "", bar,
        f"About this document. Written by Inkguard {R._version()}, code fingerprint {fp[:16]}, "
        f"for run {report.get('run_id', '')}.",
        f"Open-source software · Apache License 2.0 · {R.SOURCE.replace('https://', '')}",
        f"End of {R.CITE_NAME.lower()}.",
        "",
    ])


# =====================================================================
# 6_PLEASE-RETURN-...Feedback-report.docx
# =====================================================================

FEEDBACK_BOX = "Your feedback (optional)"
FEEDBACK_HINT = ("Anything you would like us to know. Describe the kind of problem, never the content: "
                 "\"a 7-digit number after the word 'enc' was kept\", not the number. No names, dates, "
                 "record numbers or note text.")
CONTACT_BOX = "Your contact details (optional, only if you would like a reply)"
SYSTEM_HEAD = "SYSTEM RECORD · written by Inkguard from counts and codes only · do not edit"
SYSTEM_BEGIN = "--- begin system record ---"
SYSTEM_END = "--- end system record · fingerprint {sha} ---"


def _shade(cell_or_par, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    el.append(shd)


def _borders(table, color: str, size: int = 8) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:color"), color.lstrip("#"))
        b.append(e)
    tblPr.append(b)


def _editable(paragraph, ident: int) -> None:
    """Mark a paragraph as a region anyone may edit in a protected document."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    start = OxmlElement("w:permStart")
    start.set(qn("w:id"), str(ident))
    start.set(qn("w:edGrp"), "everyone")
    end = OxmlElement("w:permEnd")
    end.set(qn("w:id"), str(ident))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def system_lines(record: list[str]) -> tuple[list[str], str]:
    """The system record's lines, and their fingerprint."""
    sha = hashlib.sha256("\n".join(record).encode("utf-8")).hexdigest()[:16]
    return record, sha


def feedback_docx(record: list[str], run_id: str, written_utc: str, path: Path) -> None:
    """The feedback report. The user's section is a blue-edged box with room
    to write and their optional contact details; it is the only part that can
    be edited (the rest is protected, with no password). Under a heavy rule,
    the system record, in grey blocks, closed by its fingerprint."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    from inkguard import cli_report as R
    rgb = (lambda h: RGBColor.from_string(h.lstrip("#").upper()))
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(2)

    def para(text="", size=9.5, bold=False, color=INK, italic=False, after=2, font=None, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        if align:
            p.alignment = align
        if text:
            r = p.add_run(text)
            r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
            r.font.color.rgb = rgb(color)
            if font:
                r.font.name = font
        return p

    title = para("", after=0)
    if LOGO.exists():
        title.add_run().add_picture(str(LOGO), width=Inches(0.17))
        title.add_run("  ")
    r = title.add_run(f"Inkguard · {R.FEEDBACK_NAME}")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, rgb(INK)
    para(f"Run {run_id or 'none'} · Version {R._version()} · Written {_when(written_utc)}",
             size=8, color=MUTED, after=6)
    para(f"Please email this report to {R.REPORT_EMAIL}.", size=12, bold=True, color=RED, after=1)
    para("The system record holds no patient information; check anything you add. Your section is optional. "
         "To check it before sending: inkguard --check-feedback <this file>",
         size=8, color=MUTED, after=8)

    # The user's section: a box with a blue edge.
    box = doc.add_table(rows=2, cols=1)
    _borders(box, ACCENT, 12)
    c1, c2 = box.rows[0].cells[0], box.rows[1].cells[0]
    _shade(c1, "#ddf4ff")
    h = c1.paragraphs[0]
    r = h.add_run(FEEDBACK_BOX)
    r.font.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), rgb(ACCENT)
    hint = c1.add_paragraph()
    r = hint.add_run(FEEDBACK_HINT)
    r.font.size, r.font.color.rgb = Pt(8), rgb(MUTED)
    c2.paragraphs[0].add_run("")
    for _ in range(7):
        c2.add_paragraph("")
    for i, p in enumerate(c2.paragraphs):
        _editable(p, 100 + i)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    contact = doc.add_table(rows=1 + len(R.CONTACT_FIELDS) + 1, cols=2)
    _borders(contact, ACCENT, 6)
    head = contact.rows[0].cells[0].merge(contact.rows[0].cells[1])
    _shade(head, "#ddf4ff")
    r = head.paragraphs[0].add_run(CONTACT_BOX)
    r.font.bold, r.font.size, r.font.color.rgb = True, Pt(10), rgb(ACCENT)
    for i, label in enumerate(R.CONTACT_FIELDS, 1):
        lc, vc = contact.rows[i].cells
        lc.width, vc.width = Inches(1.4), Inches(5.9)
        rr = lc.paragraphs[0].add_run(label.rstrip(":"))
        rr.font.size, rr.font.color.rgb = Pt(9), rgb(MUTED)
        _editable(vc.paragraphs[0], 200 + i)
    last = contact.rows[-1].cells[0].merge(contact.rows[-1].cells[1])
    rr = last.paragraphs[0].add_run(f"☐ {R.CONTACT_PERMISSION}   (replace ☐ with ☒ to agree)")
    rr.font.size = Pt(9)
    _editable(last.paragraphs[0], 300)

    # The heavy rule, then the system record.
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before, rule.paragraph_format.space_after = Pt(10), Pt(4)
    pPr = rule._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "1"), ("w:color", INK.lstrip("#"))):
        bottom.set(qn(k), v)
    bd.append(bottom)
    pPr.append(bd)
    para(SYSTEM_HEAD, size=8.5, bold=True, color=MUTED, after=3)
    lines, sha = system_lines(record)
    for text in [SYSTEM_BEGIN, *lines, SYSTEM_END.format(sha=sha)]:
        p = para(text, size=7.5, font="Consolas", color=INK if text not in (SYSTEM_BEGIN,) and
                 not text.startswith("--- end") else MUTED, after=0)
        p.paragraph_format.line_spacing = 1.0
        _shade(p, SUBTLE)

    # Footer, as on every Inkguard document.
    para("", after=6)
    para(f"About this document. Written by Inkguard {R._version()}, code fingerprint {R._fingerprint()[:16]}, "
         f"for run {run_id or 'none'}. Support codes are explained in SUPPORT_CODES.md in the source repository. Open-source software · "
         f"Apache License 2.0 · {R.SOURCE.replace('https://', '')}", size=7, color=MUTED, after=1)
    para(f"End of {R.FEEDBACK_NAME.lower()}.", size=7, italic=True, color=MUTED)

    # Protect everything but the user's section (no password: a clear
    # boundary, not a lock).
    settings = doc.settings.element
    prot = OxmlElement("w:documentProtection")
    prot.set(qn("w:edit"), "readOnly")
    prot.set(qn("w:enforcement"), "1")
    settings.append(prot)
    doc.core_properties.title = f"Inkguard feedback report, run {run_id}"
    doc.core_properties.author = "Inkguard"
    doc.save(str(path))


def docx_lines(path: Path) -> list[str]:
    """The text of a Word file as lines, in reading order: paragraphs, and
    table rows as 'Label: value'. What the feedback check reads."""
    from docx import Document
    doc = Document(str(path))
    body = doc.element.body
    out: list[str] = []
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                seen = []
                cells = []
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.append(cell._tc)
                    cells.append(cell)
                if len(cells) == 2:
                    out.append(f"{cells[0].text.strip()}: {cells[1].text.strip()}")
                else:
                    for cell in cells:
                        out.extend(p.text for p in cell.paragraphs)
    return out


def pdf_text(path: Path) -> str:
    """The text of a PDF this tool wrote (for tests and the check)."""
    import pypdf
    return "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(str(path)).pages)


def document_text(path: Path) -> str:
    """Text of any document the run writes, whatever its format."""
    path = Path(path)
    suf = path.suffix.lower()
    if suf == ".pdf":
        return pdf_text(path)
    if suf == ".docx":
        return "\n".join(docx_lines(path))
    return path.read_text(encoding="utf-8")

