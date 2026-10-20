"""
PROVENANCE LINE.

One short line at the top of every text and Word file the tool writes, so
the file says what it is wherever it is copied to:

  Inkguard 0.25.0 · Safe Harbor · INCOMPLETE REVIEW · NOT READY TO SHARE · Run 20261017-7F3K

In Word it is set small (6.5 pt), in grey, with only the status in dark
red and a hairline rule beneath it.

The line carries the review status, never only the method: a file nobody has
reviewed must never look finished because it names a method. It names no
patient, no file and no person, and it carries no fingerprint (a file cannot
hold the fingerprint of the manifest that holds its own). The run id ties it
to the manifest and the report written beside it.

Spreadsheets and CSV files are not stamped: an extra row breaks every script
that reads them. Their fingerprints are in the manifest. Flattened PDF pages
are not stamped either; the PDF's plain-text copy is.

When the tool reads a file back (a later run, the tidy-up tool), the line is
recognised and skipped, so it is never mistaken for words to check.
"""
from __future__ import annotations

import re
import secrets
import time
from pathlib import Path

PREFIX = "Inkguard "
SEP = " · "
INCOMPLETE = "INCOMPLETE REVIEW · NOT READY TO SHARE"
PENDING = "PENDING REVIEW · NOT READY TO SHARE"
GREY = (0x8A, 0x8A, 0x8A)
CRIMSON = (0xA9, 0x16, 0x36)
# The whole line, exactly as the tool writes it, and nothing else: a note
# whose own first line happens to start with "Inkguard" is never touched.
_LINE = re.compile(r"Inkguard \d+\.\d+\.\d+(?: · [^\n]+)? · Run \d{8}-[A-Z0-9]{4}")
_ALPHABET = "ABCDEFGHJKMNPQRSTVWXYZ23456789"          # no look-alikes (0/O, 1/I/L)

_METHOD = {
    "safe_harbor": "Safe Harbor",
    "limited_data_set": "Limited Data Set (still PHI, data use agreement required)",
    "date_shifted": "Dates shifted (not Safe Harbor)",
}


def new_run_id() -> str:
    """Date plus four random characters: short enough to read aloud, and
    carrying nothing about the files."""
    return time.strftime("%Y%m%d") + "-" + "".join(secrets.choice(_ALPHABET) for _ in range(4))


def line(*, version: str, profile: str, status: str, run_id: str, manifest: str = "",
         opaque: bool = False) -> str:
    """The run id points to the manifest and report beside the files, so the
    line does not need to name them. `manifest` is accepted for callers that
    pass it, but is not used."""
    method = _METHOD.get(profile, profile)
    parts = [f"Inkguard {version}", method] + (["Opaque markers"] if opaque else []) \
        + [status, f"Run {run_id}"]
    return SEP.join(parts)


def is_line(text: str) -> bool:
    return bool(_LINE.fullmatch(text.rstrip("\r\n")))


def strip(text: str) -> str:
    """The text without its provenance line, if it starts with one."""
    if not text.startswith(PREFIX):
        return text
    first, sep, rest = text.partition("\n")
    return rest if is_line(first) else text


def stamp_text(path: Path, provenance: str) -> None:
    body = path.read_text(encoding="utf-8")
    path.write_text(provenance + "\n" + strip(body), encoding="utf-8")


def stamp_docx(path: Path, provenance: str) -> None:
    """A quiet first paragraph: 6.5 pt grey, the status alone in dark red,
    a hairline rule beneath. Nothing else in the document changes."""
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    d = docx.Document(str(path))
    body = d.element.body
    first = body[0] if len(body) else None
    para = d.add_paragraph()                      # created at the end ...
    fmt = para.paragraph_format
    fmt.space_before, fmt.space_after = Pt(0), Pt(9)
    # The status is the one part in colour; everything else recedes.
    status = next((s for s in (INCOMPLETE, PENDING) if s in provenance), "")
    pieces = provenance.split(status) if status else [provenance]
    for i, text in enumerate(pieces):
        if text:
            _run(para, text, GREY, bold=False)
        if status and i == 0:
            _run(para, status, CRIMSON, bold=True)
    pPr = para._p.get_or_add_pPr()               # hairline rule beneath
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "2"), ("w:space", "4"), ("w:color", "C8C8C8")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    # Word reads paragraph settings in a fixed order: the border comes
    # before the spacing python-docx already wrote.
    spacing = pPr.find(qn("w:spacing"))
    if spacing is not None:
        spacing.addprevious(bdr)
    else:
        pPr.append(bdr)
    if first is not None:                         # ... then moved to the top
        first.addprevious(para._p)
    d.save(str(path))


def _run(para, text: str, rgb: tuple, *, bold: bool) -> None:
    from docx.shared import Pt, RGBColor
    r = para.add_run(text)
    r.font.size = Pt(6.5)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*rgb)
    r.font.name = "Arial"                         # on every Windows and Mac computer
    rPr = r._r.get_or_add_rPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    spacing = OxmlElement("w:spacing")            # a touch of tracking
    spacing.set(qn("w:val"), "6")
    size = rPr.find(qn("w:sz"))                   # Word's order: spacing before size
    if size is not None:
        size.addprevious(spacing)
    else:
        rPr.append(spacing)


def stamp(path: Path, provenance: str) -> bool:
    """Stamp one written file if its kind takes a line. Returns whether it did."""
    suf = path.suffix.lower()
    if suf in (".txt", ".text", ".md"):
        stamp_text(path, provenance)
        return True
    if suf in (".docx", ".docm"):
        stamp_docx(path, provenance)
        return True
    return False
