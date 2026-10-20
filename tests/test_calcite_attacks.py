"""A second review of document cleaning tried to get patient details past it
in every way a real file allows. Each attack
that worked is kept here, so it keeps not working.

All names and numbers are invented."""
from __future__ import annotations

import csv
import io
import re
import shutil
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))

from inkguard import calcite, cli  # noqa: E402
from inkguard.policy import evaluate  # noqa: E402

ENG = calcite.Engine(evaluate)
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
VML = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
       'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"')

needs_ocr = pytest.mark.skipif(not calcite.ocr_available(), reason="tesseract is not installed")


def _all_text(path: Path) -> str:
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            return "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
    return path.read_bytes().decode("latin-1")


def _leaks(res, needles) -> dict:
    found = {}
    for o in res.outputs:
        t = _all_text(o)
        for s in needles:
            if s.lower() in t.lower():
                found.setdefault(s, []).append(o.name)
    return found


# ---------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------

def _hostile_docx(path: Path) -> Path:
    import docx
    from docx.oxml import parse_xml
    d = docx.Document()
    d.add_paragraph("Patient: Hubert Lindqvist, MRN 4471822. Seen in clinic for chest pain.")
    p = d.add_paragraph("Formula: ")
    p._p.append(parse_xml(f'<m:oMath {M}><m:r><m:t>Kowalczyk Brandauer</m:t></m:r></m:oMath>'))
    b = d.add_paragraph("Plan: follow up.")
    b._p.insert(0, parse_xml(f'<w:bookmarkStart {W} w:id="7" w:name="Lindqvist_Hubert"/>'))
    b._p.append(parse_xml(f'<w:bookmarkEnd {W} w:id="7"/>'))
    d.add_paragraph("item one", style="List Number")
    s = d.add_paragraph("Label: ")
    s._p.append(parse_xml(
        f'<w:r {VML}><w:pict><v:shape id="s1" alt="Photo of Gwendolyn Ashcroft" o:title="Gwendolyn Ashcroft">'
        f'<v:textpath string="Bartholomew Quimby"/><v:imagedata o:title="Scan Philippa Dunmore"/>'
        f'<v:textbox><w:txbxContent><w:p><w:r><w:t>see chart</w:t></w:r></w:p></w:txbxContent>'
        f'</v:textbox></v:shape></w:pict></w:r>'))
    tmp = path.with_suffix(".t.docx")
    d.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data, name = zin.read(n), n
            if n == "word/numbering.xml":
                data = data.replace(b'w:val="%1."', b'w:val="Pt Okonkwo-Vasquez %1."', 1)
            if n == "word/styles.xml":
                data = data.replace(b'<w:name w:val="Title"/>', b'<w:name w:val="Letterhead Delacroix-Ferreira"/>', 1)
            if n.startswith("word/theme/"):
                data = re.sub(rb'<a:theme ([^>]*)name="[^"]*"', rb'<a:theme \1name="Mercy Orchard Clinic"', data,
                              count=1)
            if n == "word/document.xml":
                data = data.replace(b"<w:sectPr", f'<w:altChunk {R} r:id="rIdAC1"/><w:sectPr'.encode(), 1)
            if n == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdAC1" Type="http://schemas.'
                                    b'openxmlformats.org/officeDocument/2006/relationships/aFChunk" '
                                    b'Target="afchunk.htm"/></Relationships>')
            if n == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Default Extension="htm" ContentType="text/html"/></Types>')
                data = data.replace(b"/docProps/core.xml", b"/docProps/core2.xml")
            if n == "_rels/.rels":
                data = data.replace(b"docProps/core.xml", b"docProps/core2.xml")
            if n == "docProps/core.xml":
                name = "docProps/core2.xml"
                data = re.sub(rb"<dc:creator>[^<]*</dc:creator>",
                              b"<dc:creator>Dr. Cressida Fairbanks</dc:creator><dc:description>Consult "
                              b"for Hubert Lindqvist DOB 03/14/1951</dc:description>", data)
            zout.writestr(name, data)
        zout.writestr("word/afchunk.htm", "<html><body><p>Referral from Dr. Thaddeus Winterbottom, "
                                          "SSN 123-45-6789</p></body></html>")
        zout.writestr("word/draft_backup.xml", '<?xml version="1.0"?><x>Old draft: patient lives at '
                                               '12 Orchard Lane</x>')
    tmp.unlink()
    return path


def test_hidden_places_in_a_word_file_leave_nothing(tmp_path):
    import docx
    src = _hostile_docx(tmp_path / "hostile.docx")
    res = calcite.clean_docx(src, tmp_path / "out.docx", ENG)
    assert res.status in ("clean", "held") and res.outputs, res.reason
    needles = ["Kowalczyk", "Lindqvist", "Okonkwo", "Delacroix", "Mercy Orchard", "Winterbottom",
               "123-45-6789", "Orchard Lane", "Gwendolyn", "Bartholomew", "Philippa", "Cressida",
               "03/14/1951"]
    assert _leaks(res, needles) == {}
    assert res.sweep["found"] == 0
    docx.Document(str(res.outputs[0]))                  # still opens


# ---------------------------------------------------------------------
# Excel and CSV
# ---------------------------------------------------------------------

def test_hidden_places_in_a_workbook_leave_nothing(tmp_path):
    import openpyxl
    from openpyxl.chart import LineChart, Reference
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import NamedStyle, PatternFill
    from openpyxl.worksheet.filters import FilterColumn, Filters
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "labs"
    ws.append(["Patient", "MRN", "Glucose"])
    ws.append(["Hubert Lindqvist", 4471822, 105])
    ws.append(["Marta Quispe", 5520913, 98])
    ws["C2"].number_format = '"Pt Abernathy-Whitcombe" 0'
    ws.conditional_formatting.add("A2:A10", FormulaRule(formula=['EXACT(A2,"Zebulon Fairweather")'],
                                                        fill=PatternFill("solid", start_color="FFFF00")))
    ws.auto_filter.ref = "A1:C3"
    fc = FilterColumn(colId=0)
    fc.filters = Filters(filter=["Ignatius Pemberton"])
    ws.auto_filter.filterColumn.append(fc)
    ws.sheet_properties.codeName = "Cornelius_Vandermeer"
    ns = NamedStyle(name="Octavia Blackwood style")
    wb.add_named_style(ns)
    ws["A3"].style = ns.name
    cs = wb.create_chartsheet("Trend Lucinda Harrowgate")
    ch = LineChart()
    ch.title = "Glucose trend - Lucinda Harrowgate"
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=3), titles_from_data=True)
    cs.add_chart(ch)
    cs.sheet_state = "hidden"
    src = tmp_path / "hostile.xlsx"
    wb.save(src)
    # An escaped letter: Excel shows "Lindqvist".
    parts = {}
    with zipfile.ZipFile(src) as z:
        for n in z.namelist():
            parts[n] = z.read(n)
    sheet = next(n for n in parts if n.endswith("sheet1.xml"))
    parts[sheet] = parts[sheet].replace(b"</sheetData>", b'<row r="9"><c r="A9" t="inlineStr"><is><t>Seen '
                                        b'with Hubert _x004C_indqvist, SSN 123-45-6789.</t></is></c></row>'
                                        b'</sheetData>')
    with zipfile.ZipFile(src, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    res = calcite.clean_xlsx(src, tmp_path / "out.xlsx", ENG)
    assert res.outputs, res.reason
    needles = ["Abernathy", "Zebulon", "Ignatius", "Cornelius", "Octavia", "Lucinda", "Lindqvist",
               "_x004C_indqvist", "4471822", "Quispe", "123-45-6789"]
    assert _leaks(res, needles) == {}
    openpyxl.load_workbook(res.outputs[0])              # still opens


def test_a_csv_in_another_encoding_with_a_formula_is_cleaned_and_defused(tmp_path):
    p = tmp_path / "cp1252.csv"
    p.write_bytes("Patient,MRN,Note\r\nJosé García-Muñoz,4471822,Seen today\r\nRenée Côté,5520913,"
                  "\"=HYPERLINK(\"\"http://x.example/?q=\"\"&A2,\"\"open\"\")\"\r\n".encode("cp1252"))
    res = calcite.clean_csv(p, tmp_path / "out.csv", ENG)
    assert res.outputs, res.reason
    text = res.outputs[0].read_text(encoding="utf-8-sig")
    for s in ("García", "Garcia", "Côté", "4471822", "5520913"):
        assert s not in text
    for row in csv.reader(io.StringIO(text)):
        for cell in row:
            assert not cell.startswith(("=", "+", "-", "@")), cell


def test_broken_or_oversized_files_fail_without_stopping_anything(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Patient Hubert Lindqvist")
    d.save(tmp_path / "ok.docx")
    with zipfile.ZipFile(tmp_path / "ok.docx") as zin, zipfile.ZipFile(tmp_path / "bad.docx", "w") as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "word/document.xml":
                data = data.replace(b"</w:body>", b"</w:bodyX>")
            zout.writestr(n, data)
    assert calcite.clean_file(tmp_path / "bad.docx", tmp_path / "o.docx", ENG).status == "failed"
    big = tmp_path / "big.csv"
    big.write_text("Note\n\"" + "x" * 140_000 + "\"\n")
    assert calcite.clean_file(big, tmp_path / "o.csv", ENG).status == "failed"


# ---------------------------------------------------------------------
# PDF: text the tool reads differently from how a reader sees it
# ---------------------------------------------------------------------

def _pdf(path: Path, case: str) -> Path:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawString(72, 740, "Discharge summary. HPI: 72M with chest pain, retired schoolteacher.")
    c.drawString(72, 722, "Plan: aspirin daily; follow up in clinic in 2 weeks.")
    if case == "plain":
        c.drawString(72, 700, "Patient: Hubert Lindqvist  MRN 4471822")
    elif case == "fakebold":                  # the same line printed twice, 0.4 pt apart
        c.drawString(72, 700, "Patient: Hubert Lindqvist  MRN 4471822")
        c.drawString(72.4, 700, "Patient: Hubert Lindqvist  MRN 4471822")
    elif case == "spaced":                    # letter-spaced, as in comb fields
        t = c.beginText(72, 700)
        t.setCharSpace(6)
        t.textOut("LINDQVIST HUBERT  4471822")
        c.drawText(t)
    elif case == "rotated":
        c.saveState()
        c.translate(560, 300)
        c.rotate(90)
        c.drawString(0, 0, "Patient: Hubert Lindqvist MRN 4471822")
        c.restoreState()
    elif case == "shapes":                    # text drawn as outlines: nothing to extract
        from matplotlib.path import Path as MP
        from matplotlib.textpath import TextPath
        tp = TextPath((0, 0), "Patient: Hubert Lindqvist  DOB 03/14/1951", size=12)
        pth = c.beginPath()
        for verts, code in tp.iter_segments():
            v = [x + (72 if i % 2 == 0 else 690) for i, x in enumerate(verts)]
            if code == MP.MOVETO:
                pth.moveTo(*v)
            elif code == MP.LINETO:
                pth.lineTo(*v)
            elif code == MP.CURVE3:
                pth.curveTo(v[0], v[1], v[0], v[1], v[2], v[3])
            elif code == MP.CURVE4:
                pth.curveTo(*v)
            elif code == MP.CLOSEPOLY:
                pth.close()
        c.drawPath(pth, fill=1, stroke=0)
    c.save()
    return path


@needs_ocr
def test_a_one_line_pdf_label_and_number_is_cleaned_not_stopped(tmp_path):
    res = calcite.clean_pdf(_pdf(tmp_path / "a.pdf", "plain"), tmp_path / "o.pdf", ENG)
    assert res.status == "clean", res.reason
    assert "Patient: [NAME] MRN [MRN]" in res.outputs[1].read_text()


@needs_ocr
def test_overprinted_bold_is_read_once(tmp_path):
    res = calcite.clean_pdf(_pdf(tmp_path / "a.pdf", "fakebold"), tmp_path / "o.pdf", ENG)
    txt = res.outputs[1].read_text()
    assert "PPaa" not in txt and "Lindqvist" not in txt and "4471822" not in txt


def test_letter_spaced_names_are_held(tmp_path):
    # A line set one character at a time (comb field or wide letter-spacing)
    # is held before it is written, however wide the gaps, so a surname or
    # number spelled one glyph at a time cannot pass as loose letters.
    res = calcite.clean_pdf(_pdf(tmp_path / "a.pdf", "spaced"), tmp_path / "o.pdf", ENG)
    assert res.status == "held" and "one character at a time" in res.reason
    assert res.outputs == []


def test_sideways_text_is_held(tmp_path):
    res = calcite.clean_pdf(_pdf(tmp_path / "a.pdf", "rotated"), tmp_path / "o.pdf", ENG)
    assert res.status == "held" and "sideways" in res.reason


def test_text_drawn_as_shapes_is_held(tmp_path):
    # Glyphs drawn as vector outlines make curves that typed text and tables
    # never do; the page is held before it is written, even when the shapes
    # are too faint or small for OCR to read.
    pytest.importorskip("matplotlib")
    res = calcite.clean_pdf(_pdf(tmp_path / "a.pdf", "shapes"), tmp_path / "o.pdf", ENG)
    assert res.status == "held" and "shapes" in res.reason


def test_letter_spacing_joins_letters_but_not_table_columns():
    def w(x0, t, width=7.0):
        return {"x0": x0, "x1": x0 + width, "text": t}
    spaced = [w(0, "J"), w(13, "O"), w(26, "H"), w(39, "N"), w(66, "S"), w(79, "M"), w(92, "I")]
    assert calcite._letter_spaced_joins(spaced) == [False, True, True, True, False, True, True]
    column = [w(0, "1"), w(60, "2"), w(120, "3"), w(180, "4")]
    assert calcite._letter_spaced_joins(column) == [False, False, False, False]


def test_a_comb_line_is_detected_however_wide_the_gaps():
    def w(x0, t):
        return {"x0": x0, "x1": x0 + 6, "text": t}
    for gap in (10, 16, 24, 40):
        line = [w(i * gap, c) for i, c in enumerate("LINDQVIST")]
        assert calcite._comb_line(line), gap
    # An ordinary sentence (few single-character words) is not a comb line.
    sentence = [w(0, "I"), w(20, "saw"), w(50, "a"), w(70, "patient"), w(130, "today")]
    assert not calcite._comb_line(sentence)


def test_the_page_check_forgives_ocr_slips_but_not_names():
    read = ["Patient:", "[NAME]", "Tolerating", "diet.", "03/14/1951"]
    # unread_words is now strict: only an exact token match counts as read.
    assert calcite.unread_words(["Patient", "diet", "03/14/1951"], read) == []
    assert "Toleratinq" in calcite.unread_words(["Toleratinq"], read)
    # _page_check forgives a misread of an extracted word (an OCR slip), but
    # holds a word the tool never read. (Names drawn as shapes are held
    # earlier by the curve guard; see test_text_drawn_as_shapes_is_held.)
    assert calcite._page_check(["Toleratinq"], read, ENG) is False
    assert calcite._page_check(["Lindqvist"], read, ENG) is True
    assert calcite._page_check(["Smith", "Jones", "Brown"], read, ENG) is True


def test_without_ocr_a_pdf_is_held_not_passed(tmp_path, monkeypatch):
    monkeypatch.setattr(calcite, "ocr_available", lambda: False)
    res = calcite.clean_pdf(_pdf(tmp_path / "a.pdf", "plain"), tmp_path / "o.pdf", ENG)
    assert res.status == "held" and "OCR" in res.reason


# ---------------------------------------------------------------------
# The batch run
# ---------------------------------------------------------------------

def _docx(path: Path, text: str = "Seen today. Plan: rest.") -> Path:
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    d.save(path)
    return path


def test_one_file_that_breaks_does_not_stop_the_run(tmp_path, monkeypatch):
    inp = tmp_path / "in"
    (inp / "p1").mkdir(parents=True)
    (inp / "p2").mkdir()
    _docx(inp / "p1" / "consult.docx")
    (inp / "p2" / "note.txt").write_text("Seen today. Plan: rest.\n")
    real = calcite.clean_file

    def flaky(src, dst, eng):
        if src.name == "consult.docx":
            raise RuntimeError("disk went away")
        return real(src, dst, eng)
    monkeypatch.setattr(calcite, "clean_file", flaky)
    stats = cli.RunStats()
    cli.process_directory(evaluate, inp, tmp_path / "out", stats, None, False, True, mode="mask")
    out = {Path(e.source).name: e for e in stats.files}
    assert out["consult.docx"].outcome == "set_aside"
    assert out["note.txt"].outcome == "deidentified"
    assert stats.accounting["verified"], stats.accounting["problems"]


def test_two_files_never_share_an_output_name(tmp_path):
    inp = tmp_path / "in" / "p1"
    inp.mkdir(parents=True)
    _docx(inp / "consult.docx")
    shutil.copy(inp / "consult.docx", inp / "consult.docm")
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(inp / "visit.pdf"))
    c.drawString(72, 740, "Seen today for knee pain. Plan: rest and ice, return as needed.")
    c.save()
    (inp / "visit.pdf.txt").write_text("Seen today. Knee better.\n")   # the PDF's text copy's name
    stats = cli.RunStats()
    cli.process_directory(evaluate, tmp_path / "in", tmp_path / "out", stats, None, False, True, mode="mask")
    outs = [e.output.lower() for e in stats.files if e.outcome in ("deidentified", "held")]
    assert len(outs) == 4 and len(set(outs)) == 4, outs
    assert stats.accounting["verified"], stats.accounting["problems"]
    on_disk = [p for p in (tmp_path / "out").rglob("*") if p.is_file() and not p.name.startswith("_")]
    assert len({p.name.lower() for p in on_disk}) == len(on_disk)


def test_a_file_that_vanishes_mid_run_is_reported_not_fatal(tmp_path, monkeypatch):
    inp = tmp_path / "in" / "p1"
    inp.mkdir(parents=True)
    (inp / "a.txt").write_text("Seen today.\n")
    (inp / "b.png").write_bytes(b"not really a picture")
    real = cli._sha256

    def sha(path):
        h = real(path)
        if Path(path).name == "b.png" and Path(path).parent.name == "p1" and "in" in Path(path).parts:
            Path(path).unlink()
        return h
    monkeypatch.setattr(cli, "_sha256", sha)
    stats = cli.RunStats()
    cli.process_directory(evaluate, tmp_path / "in", tmp_path / "out", stats, None, False, True, mode="mask")
    assert not stats.accounting["verified"]
    assert any("changed during the run" in p for p in stats.accounting["problems"])
    assert (tmp_path / "out" / "3_NOT-PROCESSED-MAY-CONTAIN-PHI" / "_ORIGINAL-NAMES-KEY.csv").exists()


def test_the_sweep_finds_names_however_a_file_spells_them():
    s = {"Hubert Lindqvist", "O'Brien-Smith", "Mercy & Orchard"}
    views = {
        "surname alone": "Mr Lindqvist called", "upper": "HUBERT LINDQVIST",
        "underscore": "Lindqvist_Hubert", "amp": '<a:theme name="Mercy &amp; Orchard"/>',
        "apos": '<x v="O&apos;Brien-Smith"/>', "char ref": "<x>Hubert &#76;indqvist</x>",
    }
    for k, v in views.items():
        # Person names are also searched word by word ("Mr Lindqvist").
        assert calcite.sweep_texts(calcite._xml_views("p.xml", v.encode()), s,
                                   {"Hubert Lindqvist"})["found"] >= 1, k
    utf16 = "Hubert Lindqvist".encode("utf-16-le")
    assert calcite.sweep_texts(calcite._xml_views("p.xml", utf16), s)["found"] >= 1


def test_folder_names_holding_identifiers_are_renamed():
    n = cli._Names(evaluate)
    for name in ("Lindqvist", "4471822", "MRN4471822", "HL 03-14-1951", "LINDQVIST H"):
        assert n.flagged(name), name
