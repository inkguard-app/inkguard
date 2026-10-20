"""Found by the reliability review before the first public release:
ways an ordinary user, by accident, could make a run lose,
overwrite or mislabel a file. Each test is the smallest folder that showed
the problem. All notes are invented."""
import csv
import json

import pytest

from inkguard import cli


def _accounting(out):
    return json.loads((out / ".inkguard" / "audit_trail.json").read_text(encoding="utf-8"))["files"]


def _outcomes(out):
    m = out / cli.PILE_NOT / cli.KEY_FILE
    return {r["source_path"]: r for r in csv.DictReader(m.read_text(encoding="utf-8").splitlines())}


def test_input_files_named_like_the_tools_own_files_are_not_overwritten(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    for name in (cli.AUDIT_FILE, cli.FEEDBACK_FILE, cli.CITE_FILE, cli.NOTICE_FILE,
                 cli.WORDS_FILE, cli.KEY_FILE):
        (src / name).write_text("Patient seen for cough, doing well.\n", encoding="utf-8")
    out = tmp_path / "out"
    cli.main([str(src), "-o", str(out), "--quiet"])
    rows = _outcomes(out)
    for name, r in rows.items():
        if r["output_path"]:
            assert r["output_path"].lower() not in {n.lower() for n in cli._OURS_TOP}, r
            assert (out / r["output_path"]).read_text(encoding="utf-8").rstrip().endswith("doing well.")
    assert _accounting(out)["every_file_accounted_for"] is True


def test_a_patient_folder_named_like_a_results_folder_does_not_collide(tmp_path):
    src = tmp_path / "in"
    (src / cli.PILE_FLAGGED / "p2").mkdir(parents=True)
    (src / "p2").mkdir()
    (src / cli.PILE_FLAGGED / "p2" / "b.txt").write_text("Patient seen, doing well.\n")
    (src / "p2" / "b.txt").write_text("Patient seen. Code QX7Z9K2B noted.\n")
    out = tmp_path / "out"
    cli.main([str(src), "-o", str(out), "--quiet"])
    acc = _accounting(out)
    assert acc["every_file_accounted_for"] is True, acc["problems"]


def test_a_csv_run_refuses_to_write_over_its_input(tmp_path):
    f = tmp_path / "data.csv"
    f.write_text("id,note\n1,Patient John Smith has flu.\n", encoding="utf-8")
    before = f.read_bytes()
    with pytest.raises(SystemExit):
        cli.main([str(f), "-o", str(f), "--quiet"])
    assert f.read_bytes() == before


def test_input_inside_the_results_folder_is_refused(tmp_path):
    out = tmp_path / "out"
    src = out / cli.PILE_NOT
    (src / "pZ").mkdir(parents=True)
    (src / "pZ" / "converted.txt").write_text("Patient John Smith has flu.\n")
    with pytest.raises(SystemExit) as exc:
        cli.main([str(src), "-o", str(out), "--quiet"])
    assert "RUN-FOLDER-OVERLAP" in str(exc.value)


@pytest.mark.parametrize("encoding", ["latin-1", "cp1252", "utf-8-sig", "utf-16"])
def test_notes_saved_in_older_encodings_are_read_correctly(tmp_path, encoding):
    src = tmp_path / "in" / "p"
    src.mkdir(parents=True)
    (src / "n.txt").write_bytes(
        "Patient José Müller seen by Dr. Søren Ægir. Phone 555-555-0123.\n".encode(encoding))
    out = tmp_path / "out"
    cli.main([str(src.parent), "-o", str(out), "--quiet"])
    written = [p for p in out.rglob("*.txt") if p.name not in cli._OURS_TOP]
    assert written, "the note was not processed"
    for p in written:
        text = p.read_text(encoding="utf-8")
        assert "�" not in text
        for frag in ("Müller", "ller", "Søren", "Ægir", "555-123"):
            assert frag not in text, (encoding, frag, text)


def test_an_excel_csv_with_a_byte_order_mark_keeps_its_first_column(tmp_path):
    f = tmp_path / "data.csv"
    f.write_bytes("id,note\n1,Patient John Smith has flu.\n".encode("utf-8-sig"))
    out = tmp_path / "clean.csv"
    cli.main([str(f), "-o", str(out), "--quiet", "--text-column", "note", "--id-column", "id"])
    header = out.read_text(encoding="utf-8-sig").splitlines()[0]
    assert header.split(",")[0] == "id"
    assert "Smith" not in out.read_text(encoding="utf-8-sig")


def test_resume_cleans_a_note_again_if_it_changed(tmp_path):
    src = tmp_path / "in" / "p"
    src.mkdir(parents=True)
    note = src / "n.txt"
    note.write_text("Patient seen for cough, doing well.\n")
    out = tmp_path / "out"
    cli.main([str(src.parent), "-o", str(out), "--quiet"])
    note.write_text("Patient seen for cough. Wife Hildegard Quistorp called.\n")
    cli.main([str(src.parent), "-o", str(out), "--quiet", "--resume"])
    rows = _outcomes(out)
    r = rows["p/n.txt"]
    assert r["outcome"] != "already_present", r
    for p in out.rglob("*.txt"):
        assert "Quistorp" not in p.read_text(encoding="utf-8")


def test_resume_with_different_settings_is_refused(tmp_path):
    src = tmp_path / "in" / "p"
    src.mkdir(parents=True)
    (src / "n.txt").write_text("Patient seen 03/14/2071 for cough.\n")
    out = tmp_path / "out"
    cli.main([str(src.parent), "-o", str(out), "--quiet"])
    with pytest.raises(SystemExit) as exc:
        cli.main([str(src.parent), "-o", str(out), "--quiet", "--resume", "--profile", "lds"])
    assert "RUN-RESUME-SETTINGS" in str(exc.value)
    # The same settings resume normally.
    cli.main([str(src.parent), "-o", str(out), "--quiet", "--resume"])


def test_a_word_file_the_tool_wrote_reads_back_without_its_label(tmp_path):
    import docx
    from inkguard import lapis
    src = tmp_path / "in" / "pA"
    src.mkdir(parents=True)
    d = docx.Document()
    d.add_paragraph("Patient John Smith seen 03/02/2021 for cough.")
    d.save(str(src / "note.docx"))
    out = tmp_path / "out"
    cli.main([str(src.parent), "-o", str(out), "--quiet"])
    first = next(out.rglob("*note.docx"))
    again = tmp_path / "again" / "pA"
    again.mkdir(parents=True)
    (again / "note.docx").write_bytes(first.read_bytes())
    out2 = tmp_path / "out2"
    cli.main([str(again.parent), "-o", str(out2), "--quiet"])
    second = next(out2.rglob("*note.docx"))
    assert cli.PILE_FLAGGED not in second.relative_to(out2).as_posix()
    paras = [p.text for p in docx.Document(str(second)).paragraphs]
    assert sum(1 for p in paras if lapis.is_line(p)) == 1, paras
    assert not any(p.startswith("[NAME] 0.25") for p in paras), paras


def test_a_very_long_run_of_letters_is_set_aside_not_analysed(tmp_path):
    # 12,000 letters with no space took over 4 GB and killed the run.
    src = tmp_path / "in"
    (src / "a").mkdir(parents=True)
    (src / "z").mkdir()
    (src / "a" / "n.txt").write_text("Patient seen, doing well.\n")
    (src / "z" / "n.txt").write_text("Patient seen. " + "a" * 12000 + "\n")
    out = tmp_path / "out"
    cli.main([str(src), "-o", str(out), "--quiet"])
    rows = _outcomes(out)
    assert rows["z/n.txt"]["outcome"] == "set_aside"
    assert (out / cli.NOTICE_FILE).exists()


def test_the_typo_check_ignores_tokens_no_word_is_that_long():
    from inkguard import agate
    assert agate._typo_of_known("x" * 5000) is False
    assert agate._typo_of_known("pateint") is True


def test_the_output_manifest_vouches_only_for_files_this_run_wrote(tmp_path):
    src = tmp_path / "in"
    (src / "p1").mkdir(parents=True)
    (src / "p1" / "n.txt").write_text("Patient seen, well.\n")
    (src / "loose.txt").write_text("Patient seen, well.\n")
    out = tmp_path / "out"
    (out / "archive").mkdir(parents=True)
    (out / "archive" / "patient_list.txt").write_text("Real names: John Smith\n")
    cli.main([str(src), "-o", str(out), "--quiet"])
    rows = {r["file"]: r for r in csv.DictReader(
        (out / cli.TECH_DIR / cli.MANIFEST_FILE).read_text(encoding="utf-8").splitlines())}
    foreign = rows["archive/patient_list.txt"]
    assert foreign["status"] == cli.NOT_OURS and foreign["sha256"] == "" and foreign["run_id"] == ""
    # A loose file at the top (its name may be made neutral) is a cleaned copy too.
    loose = [r for f, r in rows.items() if f.startswith(cli.PILE_DONE + "/") and f.count("/") == 1]
    assert loose and all(r["status"] == cli.STATUS_DONE for r in loose), loose
    assert rows[cli.done_path("", "p1/n.txt").as_posix()]["status"] == cli.STATUS_DONE
    assert rows[cli.AUDIT_FILE]["status"] == "report or guide"


def test_set_aside_files_give_their_own_exit_status(tmp_path):
    src = tmp_path / "in" / "p1"
    src.mkdir(parents=True)
    (src / "n.txt").write_text("Patient seen, well.\n")
    (src / "scan.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    assert cli.main([str(src.parent), "-o", str(tmp_path / "out"), "--quiet"]) == 3
    (src / "scan.png").unlink()
    assert cli.main([str(src.parent), "-o", str(tmp_path / "out2"), "--quiet"]) == 0


def test_windows_reserved_names_are_renamed(tmp_path):
    src = tmp_path / "in" / "visits"
    src.mkdir(parents=True)
    for n in ("CON.txt", "aux.txt"):
        (src / n).write_text("Patient seen, well.\n")
    out = tmp_path / "out"
    cli.main([str(src.parent), "-o", str(out), "--quiet"])
    names = {cli._base(p.relative_to(out)).stem.lower() for p in (out / cli.PILE_DONE / "visits").iterdir()}
    assert not names & {"con", "aux"}


def test_an_output_path_that_is_a_file_gives_a_code_not_a_traceback(tmp_path):
    src = tmp_path / "in" / "p1"
    src.mkdir(parents=True)
    (src / "n.txt").write_text("Patient seen, well.\n")
    afile = tmp_path / "afile"
    afile.write_text("x")
    with pytest.raises(SystemExit) as exc:
        cli.main([str(src.parent), "-o", str(afile), "--quiet"])
    assert "RUN-OUTPUT-WRITE" in str(exc.value)


def test_a_full_disk_mid_run_gives_a_code_and_leaves_the_folder_labelled(tmp_path, monkeypatch):
    # Simulated by failing the third write of a cleaned note.
    import errno
    src = tmp_path / "in"
    for i in range(4):
        (src / f"p{i}").mkdir(parents=True)
        (src / f"p{i}" / "n.txt").write_text("Patient seen, well.\n")
    out = tmp_path / "out"
    real = cli._stamp
    calls = {"n": 0}

    def full(*a, **k):
        calls["n"] += 1
        if calls["n"] == 3:
            raise OSError(errno.ENOSPC, "No space left on device")
        return real(*a, **k)
    monkeypatch.setattr(cli, "_stamp", full)
    rc = None
    try:
        rc = cli.main([str(src), "-o", str(out), "--quiet"])
    except SystemExit as exc:
        assert "RUN-DISK-FULL" in str(exc)
    # Either the file was stopped and reported, or the run stopped with the
    # code; in both cases the folder is labelled.
    assert (out / cli.NOTICE_FILE).exists()
    if rc is not None:
        assert rc == 1


def test_an_interrupted_run_resumes_without_doubling_its_output(tmp_path, monkeypatch):
    # The run is killed after the second patient.
    src = tmp_path / "in"
    for i in range(4):
        (src / f"p{i}").mkdir(parents=True)
        (src / f"p{i}" / "n.txt").write_text("Patient seen, well.\n")
    out = tmp_path / "out"
    real = cli._write_review_list
    calls = {"n": 0}

    class Killed(BaseException):
        pass

    def die(*a, **k):
        real(*a, **k)
        calls["n"] += 1
        if calls["n"] == 2:
            cli._write_manifest(cli.local_only_dir(out), _stats["s"])
            raise Killed()
    _stats = {}
    orig = cli.process_directory

    def wrapped(evaluate, in_dir, out_dir, stats, *a, **k):
        _stats["s"] = stats
        return orig(evaluate, in_dir, out_dir, stats, *a, **k)
    monkeypatch.setattr(cli, "_write_review_list", die)
    monkeypatch.setattr(cli, "process_directory", wrapped)
    with pytest.raises(Killed):
        cli.main([str(src), "-o", str(out), "--quiet"])
    assert "DID NOT FINISH" in (out / cli.NOTICE_FILE).read_text(encoding="utf-8")
    monkeypatch.undo()
    cli.main([str(src), "-o", str(out), "--quiet", "--resume"])
    notes = sorted(p.relative_to(out).as_posix() for p in out.rglob("*__n.txt"))
    assert notes == [cli.done_path("", f"p{i}/n.txt").as_posix() for i in range(4)], notes
    assert _accounting(out)["every_file_accounted_for"] is True
    assert "DID NOT FINISH" not in (out / cli.NOTICE_FILE).read_text(encoding="utf-8")


# ---------------------------------------------- newsworthy events

@pytest.mark.parametrize("text,detail", [
    ("Brought in with two of the 14 students injured in the school bus rollover. Femur fracture.",
     "two of the 14 students"),
    ("Crush injury to the left leg in the Harlowe earthquake. Fasciotomy performed.", "Harlowe"),
    ("Gunshot wound, one of several victims of the mall shooting at the food court.", None),
    ("Injured in the Route 9 pile-up with 11 vehicles involved; her brother was also admitted.",
     "11 vehicles"),
])
def test_a_newsworthy_event_is_a_reminder_with_its_details_suggested(text, detail):
    """The details that pin an event down can go; the
    rest is a gentle reminder, not a barrier."""
    from inkguard import contract
    st = contract.sanitize(text)
    assert contract.newsworthy_events(st)
    details = [st[a:b] for _w, a, b in contract.event_details(st)]
    if detail:
        assert detail in details
    assert "earthquake" not in details and "pile-up" not in details    # the kind stays


@pytest.mark.parametrize("text", [
    "Blast cells 20% on the peripheral smear.",
    "Evacuation of the subdural hematoma via burr hole.",
    "Night terrors since age 4.",
    "MVC, restrained driver, airbag deployed, ambulatory at scene.",
    "Flooding sensation in the ears with Meniere disease.",
])
def test_ordinary_clinical_words_are_not_events(text):
    from inkguard import contract
    assert contract.newsworthy_events(text) == []


def test_an_events_name_goes_unless_both_reviewers_keep_it():
    from inkguard import chert, contract
    text = ("Injured in the Bluebird Festival stampede; crush injury to the left leg.\n"
            "Fasciotomy performed, compartment pressures normalised.\n")
    res = contract.deidentify(text, mode="mask")
    note, sanitized = chert.note_payload(text, res.raw, False)
    ev = [f for f in note["rare"] if f.get("kind") == "event"]
    assert ev and sanitized[ev[0]["start"]:ev[0]["end"]].startswith("Injured in the")
    det = [f for f in note["rare"] if f.get("kind") == "event_detail"]
    assert det and sanitized[det[0]["start"]:det[0]["end"]] == "Bluebird Festival"
    note["text"] = sanitized
    silent = {"answers": {}, "confirmed_at": "x"}
    items, _ = chert.note_items(note, silent, silent, sanitized)
    # The detail goes by default; the sentence stays unless someone says so.
    assert [it["key"] for it in items if it["entity"] == "EVENT"] == [det[0]["id"]]
    keep = {"answers": {det[0]["id"]: "fine"}, "confirmed_at": "x"}
    items, _ = chert.note_items(note, keep, keep, sanitized)
    assert not any(it["entity"] == "EVENT" for it in items)


def test_opaque_markers_reveal_nothing_about_what_was_removed():
    # Gaps no one can decode, in the FINAL files only.
    import re
    from inkguard import chert, contract
    text = ("Patient Rosalind Ferraro, MRN 0071-4482, seen 03/14/2071 by Dr. Kwame Osei. "
            "Phone (218) 555-0176. A1c 7.4%.\n")
    typed = contract.deidentify(text, mode="mask").clean_text
    assert "[NAME]" in typed and "[REMOVED]" not in typed        # drafts stay typed
    with contract.opaque_scope():
        r = contract.deidentify(text, mode="mask")
        spans = [{"start": s, "end": e, "entity": "PERSON", "layer": "review"}
                 for s, e, _ in r.edits[:1]]
        rendered = chert._render(contract.sanitize(text), spans, False)
    assert r.reproduced
    for out in (r.clean_text, rendered):
        assert "[REMOVED]" in out
        assert not re.search(r"\[(?:NAME|MRN|PHONE|DATE|ID)\]", out), out
    assert "Ferraro" not in r.clean_text and "A1c 7.4%" in r.clean_text
    assert "[NAME]" in contract.deidentify(text, mode="mask").clean_text   # scope ended


def test_drafts_never_offer_opaque_markers():
    with pytest.raises(SystemExit):
        cli.build_parser().parse_args(["in", "-o", "out", "--opaque"])


def test_a_spelling_slip_is_a_question_for_both_reviewers_but_does_not_hold_the_note():
    """Typos are recorded and seen by both reviewers and
    the doctor; they do not hold the note on their own."""
    from inkguard import chert, contract
    text = "Pt seen, abdomin soft, plan reviewed.\n"
    r = contract.deidentify(text, mode="mask")
    assert r.status == "clear"
    assert any(w["reason"] == "spelling_slip" for w in r.words_to_check)
    note, _ = chert.note_payload(text, r.raw, False)
    assert any(q["reason"] == "spelling_slip" for q in note["questions"])


def test_an_interrupted_run_is_started_over_from_the_originals(tmp_path, monkeypatch):
    """If a run is interrupted, the next run deletes what
    it wrote (copies only, never an original) and starts over."""
    src = tmp_path / "in"
    for i in range(4):
        (src / f"p{i}").mkdir(parents=True)
        (src / f"p{i}" / "n.txt").write_text("Patient seen, well.\n")
    out = tmp_path / "out"
    (out / "mine").mkdir(parents=True)
    (out / "mine" / "keep.txt").write_text("not Inkguard's")
    real = cli._write_review_list
    calls = {"n": 0}

    class Killed(BaseException):
        pass

    def die(*a, **k):
        real(*a, **k)
        calls["n"] += 1
        if calls["n"] == 2:
            raise Killed()
    monkeypatch.setattr(cli, "_write_review_list", die)
    with pytest.raises(Killed):
        cli.main([str(src), "-o", str(out), "--quiet"])
    assert cli.unfinished(out)
    assert (out / cli.TECH_DIR / cli.SNAPSHOT_DIR).exists()      # the copies being worked on
    monkeypatch.undo()
    before = {p: p.read_bytes() for p in src.rglob("*") if p.is_file()}
    cli.main([str(src), "-o", str(out), "--quiet"])
    assert {p: p.read_bytes() for p in src.rglob("*") if p.is_file()} == before
    assert not cli.unfinished(out)
    assert not (out / cli.TECH_DIR / cli.SNAPSHOT_DIR).exists()
    assert (out / "mine" / "keep.txt").read_text() == "not Inkguard's"   # only our own paths are deleted
    notes = sorted(p.relative_to(out).as_posix() for p in out.rglob("*__n.txt"))
    assert notes == [cli.done_path("", f"p{i}/n.txt").as_posix() for i in range(4)], notes
    assert _accounting(out)["every_file_accounted_for"] is True


def test_scrambled_names_are_random_codes_grouped_by_patient(tmp_path):
    import re
    src = tmp_path / "in"
    for pt in ("Garcia_Rosa", "p2"):
        (src / pt).mkdir(parents=True)
        for n in ("visit.txt", "labs.txt"):
            (src / pt / n).write_text("Patient seen, well.\n")
    out = tmp_path / "out"
    cli.main([str(src), "-o", str(out), "--quiet", "--scramble-names"])
    files = [p.relative_to(out) for p in (out / cli.PILE_DONE).rglob("*") if p.is_file()]
    assert len(files) == 4
    for f in files:
        assert re.fullmatch(r"patient-[A-Z0-9]{5}", f.parts[1]), f
        assert re.fullmatch(cli.TAG_DONE + r"__file-[A-Z0-9]{5}\.txt", f.parts[2]), f
    assert len({f.parts[1] for f in files}) == 2                 # still grouped by patient
    rows = _outcomes(out)
    assert {r["output_path"] for r in rows.values()} == {f.as_posix() for f in files}
    assert "randomised" in __import__('inkguard.cli_docs', fromlist=['x']).document_text(out / cli.AUDIT_FILE)
