"""Documents (calcite): Word, Excel, PDF and CSV files are cleaned
in every part, not only the visible text, and every clean copy is searched
for every string that was removed. The batch tool accounts for every file of
every patient; images and files it cannot read are set aside, flagged, and
never exported."""
import csv
import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ.setdefault("PHI_ROOT_KEY", "ab" * 32)

import pytest  # noqa: E402

import doc_fixtures as F  # noqa: E402
from inkguard import calcite  # noqa: E402
from inkguard.policy import evaluate  # noqa: E402

ENG = calcite.Engine(evaluate)


def everything_in(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())


# ---------------------------------------------------------------- pieces

def test_canvas_keeps_separators_and_drops_what_it_may():
    cv = calcite.Canvas()
    cv.slot("a", "Hubert")
    cv.sep(" ", droppable=True, ref="hyphen")
    cv.slot("b", "Lindqvist")
    cv.sep("\n")
    cv.slot("c", "seen")
    new, joined = cv.apply([(0, 16, "[NAME]")])
    assert new == ["[NAME]", "", "seen"] and joined == "[NAME]\nseen"
    assert cv.dropped == ["hyphen"]
    cv2 = calcite.Canvas()
    cv2.slot("a", "Rose")
    cv2.sep("\n")
    cv2.slot("b", "Alvarado")
    new, joined = cv2.apply([(0, 13, "[NAME]")])
    assert new == ["[NAME]", ""] and joined == "[NAME]\n"          # a line break stays


def test_the_sweep_is_exact_and_whole_word():
    s = calcite.sweep_texts({"a": "BP rose to 150. Rosemary called.", "b": "Seen by Rose  Alvarado."},
                            {"Rose Alvarado", "Rose", "ab"})
    assert s["found"] == 2 and s["found_in"] == ["b"] and s["strings_searched"] == 2


def test_author_names_skip_titles_and_generic_accounts():
    got = calcite._author_names(["Dr. Margaret Ellsworth", "Microsoft Office User", "jroe", "Admin"])
    assert "Margaret" in got and "Ellsworth" in got and "jroe" in got
    assert "Dr" not in got and "Admin" not in got and "Microsoft Office User" not in got


# ---------------------------------------------------------------- Word

@pytest.fixture(scope="module")
def docx_result(tmp_path_factory):
    d = tmp_path_factory.mktemp("docx")
    src = F.make_docx(d / "consult.docx")
    return src, calcite.clean_docx(src, d / "out" / "consult.docx", ENG)


def test_word_every_part_is_clean(docx_result):
    src, r = docx_result
    assert r.status == "clean", r.reason
    raw = everything_in(r.outputs[0])
    for s in F.SECRETS + ("Kellerman", "Margaret", "jroe", "MERGEFIELD", "mailto:"):
        assert s not in raw, s
    assert r.sweep["found"] == 0 and r.sweep["parts_searched"] >= 10
    found = r.hidden_found
    for k in ("comments", "tracked changes", "hidden text", "fields", "links", "text boxes",
              "headers", "footers", "document properties", "template path", "thumbnail image"):
        assert found.get(k), k


def test_word_reads_as_expected(docx_result):
    import docx
    _, r = docx_result
    d = docx.Document(str(r.outputs[0]))
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Patient [NAME] reports improvement" in body          # a name split across runs
    assert "Return in 2 weeks." in body                            # an insertion accepted
    assert "SSN" not in body                                       # a deletion gone for good
    assert "Addressed to: [NAME]" in body                          # a field became its result
    assert "[image removed]" in body
    assert d.sections[0].header.paragraphs[0].text == "Patient: [NAME]   MRN: [MRN]"
    assert d.sections[0].footer.paragraphs[0].text == "Seen by Dr. [NAME]"
    assert [[c.text for c in row.cells] for row in d.tables[0].rows] == [["Name", "[NAME]"], ["DOB", "1951"]]
    cp = d.core_properties
    assert not cp.author and not cp.title and not cp.last_modified_by and not cp.comments
    with zipfile.ZipFile(r.outputs[0]) as z:
        names = z.namelist()
        settings = z.read("word/settings.xml").decode()
    assert not any(n.startswith(("word/media/", "customXml/", "docProps/thumbnail")) for n in names)
    assert "word/comments.xml" not in names and "docProps/app.xml" not in names
    assert "attachedTemplate" not in settings and "removePersonalInformation" in settings


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice is not installed")
def test_word_clean_copy_opens_in_an_office_suite(docx_result, tmp_path):
    _, r = docx_result
    subprocess.run(["soffice", "--headless", "--convert-to", "txt:Text", str(r.outputs[0]),
                    "--outdir", str(tmp_path)], capture_output=True, timeout=120)
    txt = (tmp_path / "consult.txt").read_text(errors="ignore")
    assert "reports improvement" in txt and "Lindqvist" not in txt


def test_a_password_protected_or_old_word_file_is_refused(tmp_path):
    p = tmp_path / "locked.docx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    r = calcite.clean_file(p, tmp_path / "out" / "locked.docx", ENG)
    assert r.status == "failed" and "password" in r.reason and not r.outputs


# ---------------------------------------------------------------- Excel

@pytest.fixture(scope="module")
def xlsx_result(tmp_path_factory):
    d = tmp_path_factory.mktemp("xlsx")
    src = F.make_xlsx(d / "registry.xlsx")
    return src, calcite.clean_xlsx(src, d / "out" / "registry.xlsx", ENG)


def test_excel_every_sheet_and_part_is_clean(xlsx_result):
    import openpyxl
    _, r = xlsx_result
    assert r.status == "clean", r.reason
    raw = everything_in(r.outputs[0])
    for s in F.SECRETS + ("Marta", "Quispe", "5520913", "Kellerman", "Hickory", "PatientName"):
        assert s not in raw, s
    wb = openpyxl.load_workbook(r.outputs[0])
    assert all(ws.sheet_state == "visible" for ws in wb.worksheets)
    ws = wb.worksheets[0]
    assert ws.title == "(NAME) labs"
    assert not ws.row_dimensions[3].hidden
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][:3] == ("[NAME]", "[MRN]", "1951") and rows[1][4] == 105
    assert rows[2][2] == "[DATE]"                                   # born 90+ years ago
    assert all(c.comment is None for row in ws.iter_rows() for c in row)
    assert not wb.defined_names and not ws.data_validations.dataValidation
    assert not (wb.properties.creator or "") and not wb.properties.title
    assert list(wb.worksheets[1].iter_rows(values_only=True))[1] == ("[SSN]", "[NAME]")
    for k in ("hidden sheets", "hidden rows and columns", "cell notes", "defined names",
              "data validations", "headers and footers", "document properties"):
        assert r.hidden_found.get(k), k
    assert r.removed.get("formulas turned into values") == 1


# ---------------------------------------------------------------- PDF

@pytest.fixture(scope="module")
def pdf_result(tmp_path_factory):
    d = tmp_path_factory.mktemp("pdf")
    src = F.make_pdf(d / "discharge.pdf")
    return src, calcite.clean_pdf(src, d / "out" / "discharge.pdf", ENG)


def test_pdf_is_flattened_and_its_text_copy_is_clean(pdf_result):
    import pdfplumber
    _, r = pdf_result
    assert r.status == "clean", r.reason
    pdf, txt = r.outputs
    with pdfplumber.open(str(pdf)) as p:
        assert len(p.pages) == 2
        assert not "".join(pg.extract_text() or "" for pg in p.pages).strip()   # no text to recover
    text = txt.read_text()
    assert "Patient: [NAME] MRN: [MRN] DOB: 1951" in text and "Tolerating diet" in text
    for s in F.SECRETS:
        assert s not in text and s.encode() not in pdf.read_bytes(), s
    assert r.hidden_found.get("document properties")


def test_a_scanned_pdf_is_refused(tmp_path):
    from PIL import Image
    p = tmp_path / "scan.pdf"
    Image.new("RGB", (850, 1100), (250, 250, 250)).save(p, "PDF", resolution=100)
    r = calcite.clean_file(p, tmp_path / "out" / "scan.pdf", ENG)
    assert r.status == "failed" and ("scanned" in r.reason or "no text" in r.reason)


# ---------------------------------------------------------------- CSV

def test_every_column_of_a_csv_is_read(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("Name,MRN,DOB,Diagnosis\nHubert Lindqvist,4471822,03/14/1951,chest pain\n"
                 "Marta Quispe,5520913,01/02/1933,knee pain\n")
    r = calcite.clean_file(p, tmp_path / "out" / "t.csv", ENG)
    assert r.status == "clean", r.reason
    rows = list(csv.reader(r.outputs[0].open()))
    assert rows[0] == ["Name", "MRN", "DOB", "Diagnosis"]
    assert rows[1] == ["[NAME]", "[MRN]", "1951", "chest pain"]
    assert rows[2][0] == "[NAME]" and rows[2][3] == "knee pain"


# ---------------------------------------------------------------- the batch tool: every file

@pytest.fixture(scope="module")
def batch_run(tmp_path_factory):
    from inkguard import cli
    d = tmp_path_factory.mktemp("batch")
    a = d / "in" / "lindqvist_hubert"
    a.mkdir(parents=True)
    F.make_docx(a / "consult_2021-03-02.docx")
    F.make_pdf(a / "discharge.pdf")
    (a / "note.txt").write_text("PROGRESS NOTE\nHubert tolerating diet.\n")
    from PIL import Image
    Image.new("RGB", (40, 40)).save(a / "photo_id.jpg")
    (a / ".DS_Store").write_bytes(b"\x00\x00")
    (a / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0 old")
    (a / "locked.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    b = d / "in" / "p2"
    b.mkdir()
    (b / "visit.txt").write_text("Follow-up visit. Wound clean. Return as needed.\n")
    rc = cli.main([str(d / "in"), "-o", str(d / "out"), "-q"])
    return d, rc


def test_every_file_of_every_patient_is_accounted_for(batch_run):
    d, rc = batch_run
    out, local = d / "out", d / "out" / "3_NOT-PROCESSED-MAY-CONTAIN-PHI"
    rows = list(csv.DictReader((local / "_ORIGINAL-NAMES-KEY.csv").open()))
    assert len(rows) == 8
    by = {Path(r["source_path"]).name: r for r in rows}
    assert by["photo_id.jpg"]["outcome"] == "set_aside" and "image" in by["photo_id.jpg"]["reason"]
    assert by["old.doc"]["outcome"] == "set_aside"
    assert by["locked.docx"]["outcome"] == "set_aside" and "password" in by["locked.docx"]["reason"]
    assert by[".DS_Store"]["outcome"] == "system"
    for n in ("consult_2021-03-02.docx", "discharge.pdf", "note.txt", "visit.txt"):
        assert by[n]["outcome"] == "redacted", n
        assert by[n]["output_folder"] == "1_REDACTED-NO-FLAGS" and (out / by[n]["output_path"]).exists()
    # Copies of what was not processed are byte-identical, tagged, and in
    # the third folder, with their original names and folders.
    assert by["photo_id.jpg"]["output_folder"] == "3_NOT-PROCESSED-MAY-CONTAIN-PHI"
    img = out / by["photo_id.jpg"]["output_path"]
    assert img.name == "NOT-PROCESSED-MAY-CONTAIN-PHI__photo_id.jpg"
    assert img.read_bytes() == (d / "in" / "lindqvist_hubert" / "photo_id.jpg").read_bytes()
    assert (out / by[".DS_Store"]["output_path"]).exists()          # every file comes out
    assert not (local / "_ABOUT-THIS-FOLDER.txt").exists()
    rep = json.loads((out / ".inkguard" / "audit_trail.json").read_text())
    f = rep["files"]
    assert f["every_file_accounted_for"] and f["found"] == 8 and f["set_aside_not_deidentified"] == 3
    assert [p["files"] for p in f["per_patient"]] == [7, 1]
    assert rc == 1                              # a file the tool could not open is a failure


def test_output_names_carry_no_patient_information(batch_run):
    d, _ = batch_run
    from inkguard import cli
    out = d / "out"
    clean = [p for p in out.rglob("*") if p.is_file() and "3_NOT-PROCESSED-MAY-CONTAIN-PHI" not in p.parts]
    for p in clean:
        rel = str(p.relative_to(out))
        assert "lindqvist" not in rel.lower() and "2021-03-02" not in rel, rel
    md = " ".join(__import__('inkguard.cli_docs', fromlist=['x']).document_text(out / cli.AUDIT_FILE).split())
    assert "lindqvist" not in md.lower()
    assert "FILES BY FOLDER" in md


def test_a_missing_output_is_caught(tmp_path):
    from inkguard import cli
    (tmp_path / "in" / "p1").mkdir(parents=True)
    (tmp_path / "in" / "p1" / "a.txt").write_text("Knee pain, improving.\n")
    stats = cli.RunStats()
    cli.process_directory(evaluate, tmp_path / "in", tmp_path / "out", stats, None, False, True, mode="mask")
    assert stats.accounting["verified"]
    cli.done_path(tmp_path / "out", "p1/a.txt").unlink()
    cli._verify_accounting(tmp_path / "in", tmp_path / "out", stats)
    assert not stats.accounting["verified"] and "missing" in stats.accounting["problems"][0]
    (tmp_path / "in" / "p1" / "b.txt").write_text("new")
    cli._verify_accounting(tmp_path / "in", tmp_path / "out", stats)
    assert any("changed during the run" in p for p in stats.accounting["problems"])


def test_the_pre_upload_check_refuses_set_aside_files_and_the_manifest():
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
    import preflight_check as pc
    for f in ("out/unprocessed/p1/scan.jpg", "out/file_manifest.csv", "out/.work-abc/x.docx",
              "out/3_NOT-PROCESSED-MAY-CONTAIN-PHI/p1/NOT-PROCESSED-MAY-CONTAIN-PHI__scan.jpg",
              "out/3_NOT-PROCESSED-MAY-CONTAIN-PHI/_ORIGINAL-NAMES-KEY.csv"):
        assert pc.ARTIFACT_FILES.search(f), f


def test_the_three_person_release_lists_every_file_it_does_not_read(tmp_path):
    from inkguard import chert
    (tmp_path / "n" / "p1").mkdir(parents=True)
    (tmp_path / "n" / "p1" / "a.txt").write_text("Knee pain, improving. Return as needed.\n")
    from PIL import Image
    Image.new("RGB", (20, 20)).save(tmp_path / "n" / "p1" / "scan.png")
    (tmp_path / "n" / "p2").mkdir()
    Image.new("RGB", (20, 20)).save(tmp_path / "n" / "p2" / "only_a_photo.jpg")
    d = chert.prepare(tmp_path / "n", tmp_path / "b", authorized_by="Dr. O",
                      reviewers=["Jane Roe", "Sam Poe"])[0]
    b = chert.load_batch(d)
    srcs = sorted(x["source"] for x in b["not_included"])
    assert srcs == ["p1/scan.png", "p2/only_a_photo.jpg"]
    assert all("image" in x["reason"] for x in b["not_included"])
